import os
from io import BytesIO
from datetime import datetime

from flask import Flask, render_template, send_file, request
from dotenv import load_dotenv
import psycopg2
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, Border, Side, PatternFill


load_dotenv()

app = Flask(__name__)


def get_connection():
    return psycopg2.connect(
        host=os.getenv("DB_HOST"),
        port=os.getenv("DB_PORT", "5432"),
        dbname=os.getenv("DB_NAME"),
        user=os.getenv("DB_USER"),
        password=os.getenv("DB_PASSWORD")
    )


@app.route("/")
def index():

    conn = get_connection()
    cur = conn.cursor()

    cur.execute("""
        SELECT nombre
        FROM eventos_permitidos
        WHERE activo = TRUE
        ORDER BY nombre;
    """)

    eventos = [row[0] for row in cur.fetchall()]

    cur.close()
    conn.close()

    return render_template("index.html", eventos=eventos)


@app.route("/descargar")
def descargar_excel():

    evento = request.args.get("evento")

    if not evento:
        return "Debe seleccionar un evento", 400

    conn = get_connection()
    cur = conn.cursor()

    cur.execute("""
        SELECT nombre, identidad, telefono
        FROM registros_eventos
        WHERE LOWER(evento) = LOWER(%s)
        ORDER BY nombre;
    """, (evento,))

    registros = cur.fetchall()

    cur.close()
    conn.close()

    wb = Workbook()
    wb.remove(wb.active)

    # =========================
    # FECHA EN ESPAÑOL
    # =========================

    dias = {
        "Monday": "Lunes",
        "Tuesday": "Martes",
        "Wednesday": "Miércoles",
        "Thursday": "Jueves",
        "Friday": "Viernes",
        "Saturday": "Sábado",
        "Sunday": "Domingo"
    }

    meses = {
        1: "Enero",
        2: "Febrero",
        3: "Marzo",
        4: "Abril",
        5: "Mayo",
        6: "Junio",
        7: "Julio",
        8: "Agosto",
        9: "Septiembre",
        10: "Octubre",
        11: "Noviembre",
        12: "Diciembre"
    }

    ahora = datetime.now()

    dia_semana = dias[ahora.strftime("%A")]
    dia = ahora.day
    mes = meses[ahora.month]
    anio = ahora.year

    fecha_actual = f"{dia_semana} {dia} de {mes} del {anio}"

    # Nombre archivo DDMMYY
    fecha_archivo = ahora.strftime("%d%m%y")

    evento_titulo = evento.capitalize()

    headers = [
        "N°",
        "Nombre",
        "DNI",
        "N° Teléfono",
        "Firma"
    ]

    # =========================
    # ESTILOS
    # =========================

    blue_fill = PatternFill(
        start_color="1E88E5",
        end_color="1E88E5",
        fill_type="solid"
    )

    bold = Font(
        bold=True
    )

    title_font = Font(
        bold=True,
        size=14,
        color="FFFFFF"
    )

    header_font = Font(
        bold=True,
        color="FFFFFF"
    )

    center = Alignment(
        horizontal="center",
        vertical="center"
    )

    left_center = Alignment(
        horizontal="left",
        vertical="center"
    )

    border = Border(
        left=Side(style="thin"),
        right=Side(style="thin"),
        top=Side(style="thin"),
        bottom=Side(style="thin")
    )

    registros_por_hoja = 20

    partes = [
        registros[i:i + registros_por_hoja]
        for i in range(0, len(registros), registros_por_hoja)
    ]

    if not partes:
        partes = [[]]

    # =========================
    # CREAR HOJAS
    # =========================

    for hoja_num, parte in enumerate(partes, start=1):

        ws = wb.create_sheet(
            title=f"{evento[:25]} {hoja_num}"
        )

        # =========================
        # TITULO
        # =========================

        ws.merge_cells("A1:E1")

        ws["A1"] = f"Listado para {evento_titulo}"

        ws["A1"].font = title_font
        ws["A1"].alignment = center
        ws["A1"].fill = blue_fill

        # =========================
        # MOTIVO
        # =========================

        ws.merge_cells("A2:E2")

        ws["A2"] = "Motivo:"

        ws["A2"].font = bold
        ws["A2"].alignment = left_center

        # =========================
        # LUGAR
        # =========================

        ws.merge_cells("A3:E3")

        ws["A3"] = (
            "Lugar: San Vicente Centenario, "
            "Depto. Santa Bárbara"
        )

        ws["A3"].font = bold
        ws["A3"].alignment = left_center

        # =========================
        # FECHA
        # =========================

        ws.merge_cells("A4:E4")

        ws["A4"] = fecha_actual

        ws["A4"].font = bold
        ws["A4"].alignment = left_center

        # =========================
        # HEADERS
        # =========================

        ws.append(headers)

        for cell in ws[5]:

            cell.font = header_font
            cell.alignment = center
            cell.border = border
            cell.fill = blue_fill

        # =========================
        # DATOS
        # =========================

        for i, registro in enumerate(
            parte,
            start=(hoja_num - 1) * registros_por_hoja + 1
        ):

            nombre, identidad, telefono = registro

            ws.append([
                i,
                nombre,
                identidad,
                telefono,
                ""
            ])

        # =========================
        # ESTILOS TABLA
        # =========================

        for row in ws.iter_rows(
            min_row=6,
            max_row=ws.max_row,
            min_col=1,
            max_col=5
        ):

            for cell in row:

                cell.border = border
                cell.alignment = left_center

        # =========================
        # ANCHO COLUMNAS
        # =========================

        ws.column_dimensions["A"].width = 8
        ws.column_dimensions["B"].width = 42
        ws.column_dimensions["C"].width = 22
        ws.column_dimensions["D"].width = 18
        ws.column_dimensions["E"].width = 35

        ws.row_dimensions[1].height = 26

    # =========================
    # GENERAR ARCHIVO
    # =========================

    output = BytesIO()

    wb.save(output)

    output.seek(0)

    filename = f"{evento}_registros_{fecha_archivo}.xlsx"

    return send_file(
        output,
        as_attachment=True,
        download_name=filename,
        mimetype=(
            "application/"
            "vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        )
    )


if __name__ == "__main__":

    app.run(
        host="0.0.0.0",
        port=int(os.environ.get("PORT", 5000))
    )
